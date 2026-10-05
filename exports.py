from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable, List

import pandas as pd
from openpyxl import Workbook


def export_workbook(state: Dict[str, Any], output_path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "RQ Summary"
    header = ["Research Question", "Lens", "Number of Quotes", "Number of Participants", "Main Themes", "Preliminary Synthesis", "Qualitative Analytic Focus", "Participant IDs"]
    ws.append(header)
    rq_summary = state.get("rq_summary", [])
    for entry in rq_summary:
        ws.append([
            entry.get("Research Question", ""),
            entry.get("Lens", ""),
            entry.get("Number of Quotes", 0),
            entry.get("Number of Participants", 0),
            " | ".join(entry.get("Main Themes", [])),
            entry.get("Preliminary Synthesis", ""),
            entry.get("Qualitative Analytic Focus", ""),
            ", ".join(entry.get("Participant IDs", [])),
        ])

    ws = wb.create_sheet("Master Coding")
    ws.append(["Theme", "Subtheme", "Broader Insight Category", "Exact Quote", "Timestamp", "Participant ID", "Condition Code", "Topic", "Related Research Question(s)", "Note", "Code ID", "Excerpt ID", "Protocol Item(s)", "Source File", "Source Paragraphs", "Evidence Type", "Human Review", "RQ2a-1 Included", "RQ2a-2 Included", "RQ2a-3 Included", "RQ2a-4 Included"])
    for row in state.get("coding", []):
        ws.append([
            row.get("theme", ""),
            row.get("subtheme", ""),
            row.get("broader_insight_category", ""),
            row.get("exact_quote", ""),
            row.get("timestamp", ""),
            row.get("participant_id", ""),
            row.get("condition_code", ""),
            row.get("topic", ""),
            row.get("related_rqs", ""),
            row.get("note", ""),
            row.get("code_id", ""),
            row.get("excerpt_id", ""),
            row.get("protocol_items", ""),
            row.get("source_file", ""),
            row.get("source_paragraphs", ""),
            row.get("evidence_type", ""),
            row.get("human_review", ""),
            row.get("rq2a_1_included", 0),
            row.get("rq2a_2_included", 0),
            row.get("rq2a_3_included", 0),
            row.get("rq2a_4_included", 0),
        ])

    ws = wb.create_sheet("Condition Comparison")
    ws.append(["Condition", "Interview n", "Participant IDs", "Completion and Engagement", "Habit Evidence", "Routine Integration", "Reminder Experience", "Personalization and Design Preferences", "Interpretive Limits"])
    for row in state.get("condition_comparison", []):
        ws.append([
            row.get("Condition", ""),
            row.get("Interview n", ""),
            row.get("Participant IDs", ""),
            row.get("Completion and Engagement", ""),
            row.get("Habit Evidence", ""),
            row.get("Routine Integration", ""),
            row.get("Reminder Experience", ""),
            row.get("Personalization and Design Preferences", ""),
            row.get("Interpretive Limits", ""),
        ])

    ws = wb.create_sheet("Case Matrix")
    ws.append(["Participant ID", "Group", "Assigned Reminder Condition", "Assigned Course Order", "Reported Reminder Exposure", "Session Trigger and Routine", "Completion and Engagement", "Habit Interpretation", "Topic and Motivation", "Friction and Intrusiveness", "Design Preferences", "Evidence Locator"])
    for row in state.get("case_matrix", []):
        ws.append([
            row.get("Participant ID", ""),
            row.get("Group", ""),
            row.get("Assigned Reminder Condition", ""),
            row.get("Assigned Course Order", ""),
            row.get("Reported Reminder Exposure", ""),
            row.get("Session Trigger and Routine", ""),
            row.get("Completion and Engagement", ""),
            row.get("Habit Interpretation", ""),
            row.get("Topic and Motivation", ""),
            row.get("Friction and Intrusiveness", ""),
            row.get("Design Preferences", ""),
            row.get("Evidence Locator", ""),
        ])

    ws = wb.create_sheet("Codebook")
    ws.append(["Code ID", "Code", "Related Preliminary Theme", "Category", "Include", "Exclude / Interpretive Boundary", "Excerpt Records", "Example Excerpt ID"])
    for row in state.get("codebook", []):
        ws.append([
            row.get("code_id", ""),
            row.get("name", ""),
            row.get("related_theme", ""),
            row.get("category", ""),
            row.get("include", ""),
            row.get("exclude", ""),
            row.get("excerpt_records", ""),
            row.get("example_excerpt_id", ""),
        ])

    ws = wb.create_sheet("Method and Review")
    ws.append(["Aspect", "Analysis Decision / Review Requirement"])
    ws.append(["Analysis status", state.get("analysis_status", "")])
    ws.append(["Method", "Framework Method with RQ-led organizing domains and inductive codes developed from participant accounts."])
    ws.append(["Coverage", "Preliminary AI-assisted framework analysis pending human review."])

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def export_coding_csv(state: Dict[str, Any], output_path: str | Path) -> Path:
    df = pd.DataFrame(state.get("coding", []))
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)
    return path


def export_json_backup(state: Dict[str, Any], output_path: str | Path) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        json.dump(state, fh, ensure_ascii=False, indent=2)
    return path


def export_rq_report(state: Dict[str, Any], output_path: str | Path) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    report_lines = [
        "# SPARK Nudging Interview Analysis Report",
        "",
        "This report is provisional and requires researcher review before being treated as final analysis.",
        "",
    ]
    for rq in state.get("rq_summary", []):
        report_lines.append(f"## {rq.get('RQ Code', '')} - {rq.get('Lens', '')}")
        report_lines.append(rq.get("Research Question", ""))
        report_lines.append("")
        report_lines.append("Main themes:")
        for m in rq.get("Main Themes", []):
            report_lines.append(f"- {m}")
        report_lines.append("")
        report_lines.append(rq.get("Preliminary Synthesis", ""))
        report_lines.append("")
    path.write_text("\n".join(report_lines), encoding="utf-8")
    return path
